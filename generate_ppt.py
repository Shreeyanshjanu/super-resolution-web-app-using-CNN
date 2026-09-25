#!/usr/bin/env python3
"""
Generate PowerPoint presentation for Deep Learning Based Super-Resolution Mapping
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import re

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    slides_content = parse_content()
    
    for slide_data in slides_content:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        # Set title
        title_shape = slide.shapes.title
        title_shape.text = slide_data['title']
        
        # Set body content
        body_shape = slide.placeholders[1]
        text_frame = body_shape.text_frame
        text_frame.clear()
        
        for line in slide_data['content']:
            p = text_frame.add_paragraph()
            p.text = line
            p.level = 0
            p.space_after = Pt(6)
    
    # Save
    output_path = r'C:\Users\arthu\Desktop\satellite-srm\Deep_Learning_Super_Resolution.pptx'
    prs.save(output_path)
    print(f"Presentation created successfully at: {output_path}")

def parse_content():
    # Read the file
    with open(r'C:\Users\arthu\.qoder\tmp\C--Users-arthu-Desktop-satellite-srm\attachments\7d1a0a2a-b55b-484a-b87a-ee9d3ba5a75a\6fba9bee-fda1-4d9c-917f-dee7e0842ae3.txt', 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    slides = []
    current_slide = None
    
    for line in lines:
        line = line.strip()
        
        # Detect slide headers (Slide N —)
        slide_match = re.match(r'Slide\s+(\d+)\s*—\s*(.+)', line)
        if slide_match:
            if current_slide:
                slides.append(current_slide)
            
            slide_num = int(slide_match.group(1))
            slide_title = slide_match.group(2)
            current_slide = {
                'number': slide_num,
                'title': slide_title,
                'content': []
            }
        elif current_slide and line:
            current_slide['content'].append(line)
    
    if current_slide:
        slides.append(current_slide)
    
    return slides

if __name__ == '__main__':
    create_presentation()
